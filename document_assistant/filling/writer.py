"""Запись принятых ответов в копию шаблона клиента — строго по ID."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from document_assistant.filling.models import ParsedTemplate, RowAnswer
from document_assistant.reports.writers import _HEADER, ExcelReportWriter, _status_fill

NOT_PROCESSED = "Не обработано"
_GREY = "E7E6E6"

_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill("solid", fgColor=_HEADER)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_WRAP = Alignment(vertical="top", wrap_text=True)
_THIN = Side(style="thin", color="CCCCCC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# ID последним, самым правым: по нему строка шаблона находится на листе «Результат» и обратно.
_ANNOTATION_HEADERS = ["Покрытие по программе", "Статус", "Комментарий", "ID"]
_ANNOTATION_WIDTHS = [45, 14, 50, 9]
_RESULT_HEADERS = ["ID", "Лист", "Строка", "Требование клиента",
                   "Покрытие по программе", "Статус", "Комментарий"]
_RESULT_WIDTHS = [9, 18, 8, 50, 45, 14, 50]


def _fill_for(status: str) -> PatternFill:
    color = _GREY if status == NOT_PROCESSED else _status_fill(status)
    return PatternFill("solid", fgColor=color)


class TemplateAnswerWriter:
    """Колонки ответа — после последнего заполненного столбца листа, как и
    в старом пути: свои данные клиента (например, «комментарии ВСК») не трогаем.

    Каждая строка с ID получает либо принятый ответ, либо «Не обработано».
    Лист «Результат» — все строки с ID в порядке шаблона, с адресом строки:
    по нему видно, куда лёг каждый ответ.
    """

    def write(
        self,
        source: Path,
        output: Path,
        template: ParsedTemplate,
        accepted: dict[str, RowAnswer],
        summary: str,
    ) -> Path:
        shutil.copy2(source, output)
        wb = load_workbook(output)

        sheets = {r.sheet for r in template.rows}
        ann_cols = {name: ExcelReportWriter._last_used_column(wb[name]) + 1 for name in sheets}

        for name in sheets:
            ws, col = wb[name], ann_cols[name]
            for i, title in enumerate(_ANNOTATION_HEADERS):
                cell = ws.cell(row=1, column=col + i)
                if isinstance(cell, MergedCell):
                    continue
                cell.value = title
                cell.font, cell.fill, cell.alignment = _HEADER_FONT, _HEADER_FILL, _CENTER
                ws.column_dimensions[cell.column_letter].width = _ANNOTATION_WIDTHS[i]

        result_rows = []
        for row in template.rows:
            ans = accepted.get(row.row_id)
            coverage, status, comment = (
                (ans.coverage, ans.status, ans.comment) if ans else ("", NOT_PROCESSED, "")
            )
            ws, col = wb[row.sheet], ann_cols[row.sheet]
            for offset, value in enumerate((coverage, status, comment, row.row_id)):
                cell = ws.cell(row=row.row, column=col + offset)
                if isinstance(cell, MergedCell):
                    continue
                cell.value = value
                cell.alignment, cell.border = _WRAP, _BORDER
                if offset == 1:
                    cell.fill = _fill_for(status)
            result_rows.append([row.row_id, row.sheet, row.row, row.text, coverage, status, comment])

        self._write_result_sheet(wb, result_rows)
        if summary:
            ExcelReportWriter._write_summary_sheet(wb, summary)
        wb.save(output)
        return output

    @staticmethod
    def _write_result_sheet(wb, rows: list[list]) -> None:
        title = ExcelReportWriter._unique_title(wb, "Результат")
        ws = wb.create_sheet(title=title, index=0)
        for col, (name, width) in enumerate(zip(_RESULT_HEADERS, _RESULT_WIDTHS), start=1):
            cell = ws.cell(row=1, column=col, value=name)
            cell.font, cell.fill, cell.alignment = _HEADER_FONT, _HEADER_FILL, _CENTER
            ws.column_dimensions[cell.column_letter].width = width
        for r, values in enumerate(rows, start=2):
            for c, value in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c, value=value)
                cell.alignment, cell.border = _WRAP, _BORDER
                if c == 6:
                    cell.fill = _fill_for(value)
        ws.freeze_panes = "A2"
        wb.active = wb.sheetnames.index(title)
