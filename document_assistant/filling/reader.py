"""Шаблон клиента → строки с ID и таблицы для модели."""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from document_assistant.filling.models import (
    ParsedTemplate,
    SheetTable,
    TableLine,
    TemplateRow,
    format_row_id,
)
from document_assistant.reports.writers import ExcelReportWriter

logger = logging.getLogger(__name__)

# Шапка таблицы: короткие подписи, хотя бы в стольких столбцах.
_HEADER_MIN_CELLS = 3
_HEADER_MAX_LEN = 60
# Сколько строк ниже проверять: подпись столбца не повторяется в данных.
_HEADER_LOOKAHEAD = 30
# Сколько строк перед частью повторять как контекст (заголовок раздела и т.п.).
_CONTEXT_BEFORE = 3


def _cell_text(value) -> str:
    if value is None:
        return ""
    # «|» и переносы ломают markdown-таблицу, которую видит модель.
    return " ".join(str(value).replace("|", "/").split())


class TemplateReader:
    """Читает .xlsx и выдаёт строки, на которые нужен ответ.

    - Скрытые листы и столбцы пропускаются: пользователь их не видит, а модель
      на скрытых листах (в КСК — списки клиник) только тратит токены.
      Скрытые строки НЕ пропускаются: чаще всего их скрыл автофильтр, и
      пропуск молча потерял бы требования.
    - Читаются значения ячеек, а не формулы (``data_only``): иначе модель
      видит ``='[1]О Страховщике'!C4``.
    - Столбец требований определяется так же, как в старом пути
      (``ExcelReportWriter._requirement_column``).
    - Шапка таблицы может стоять не в 1-й строке (в Е-СОМ — в 11-й, под
      блоком инструкций). Она не получает ID, а её подписи становятся
      названиями столбцов в каждой части.
    - ID получает каждая строка ниже шапки (и не выше 2-й), где в столбце
      требований есть текст. Остальные непустые строки идут в таблицу без
      ID — как контекст (инструкции, разделы, детализация).
    """

    def read(self, path: str | Path) -> ParsedTemplate:
        wb = load_workbook(path, data_only=True)
        rows: list[TemplateRow] = []
        tables: list[SheetTable] = []
        counter = 0

        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                logger.info(f"Лист {ws.title!r} скрыт — пропущен")
                continue

            hidden_cols = ExcelReportWriter._hidden_columns(ws)
            last_col = ExcelReportWriter._last_used_column(ws)
            if last_col == 0:
                continue
            req_col = ExcelReportWriter._requirement_column(ws)
            if req_col in hidden_cols:
                continue
            columns = [c for c in range(1, last_col + 1) if c not in hidden_cols]
            letters = [get_column_letter(c) for c in columns]

            grid = {
                r: [_cell_text(ws.cell(row=r, column=c).value) for c in columns]
                for r in range(1, ws.max_row + 1)
            }
            req_idx = columns.index(req_col)
            header_row = self._find_header_row(grid, req_idx)

            table = SheetTable(
                sheet=ws.title,
                columns=letters,
                requirement_column=get_column_letter(req_col),
                titles=(
                    {l: t for l, t in zip(letters, grid[header_row]) if t} if header_row else {}
                ),
            )
            sheet_ids = 0
            for r, cells in grid.items():
                if not any(cells) or r == header_row:
                    continue
                req_text = cells[req_idx]
                row_id = None
                if req_text and r > (header_row or 1):
                    counter += 1
                    sheet_ids += 1
                    row_id = format_row_id(counter)
                    rows.append(TemplateRow(row_id=row_id, sheet=ws.title, row=r, text=req_text))
                table.lines.append(TableLine(row_id=row_id, cells=cells))

            if sheet_ids:
                tables.append(table)
                logger.info(
                    f"Лист {ws.title!r}: столбец требований {table.requirement_column}, "
                    f"шапка {'в строке ' + str(header_row) if header_row else 'не найдена'}, "
                    f"строк с ID {sheet_ids}, скрытых столбцов {len(hidden_cols)}")

        return ParsedTemplate(rows=rows, tables=tables)

    @staticmethod
    def _find_header_row(grid: dict[int, list[str]], req_idx: int) -> int | None:
        """Строка шапки: первая строка с текстом в столбце требований, если она
        выглядит как подписи столбцов, иначе 1-я строка, если выглядит так же.

        Подписи — это не меньше ``_HEADER_MIN_CELLS`` коротких ячеек, ни одна из
        которых не повторяется в своём столбце ниже: у строк данных повторы
        есть («обязательный минимум», название раздела), у шапки — нет.
        """
        def looks_like_header(r: int) -> bool:
            cells = grid[r]
            filled = [(i, v) for i, v in enumerate(cells) if v]
            if len(filled) < _HEADER_MIN_CELLS or any(len(v) > _HEADER_MAX_LEN for _, v in filled):
                return False
            below = [grid[x] for x in range(r + 1, r + 1 + _HEADER_LOOKAHEAD) if x in grid]
            return all(v.lower() != row[i].lower() for i, v in filled for row in below)

        first_req = next((r for r, cells in grid.items() if r >= 2 and cells[req_idx]), None)
        if first_req and looks_like_header(first_req):
            return first_req
        if 1 in grid and looks_like_header(1):
            return 1
        return None

    @staticmethod
    def chunks(template: ParsedTemplate, batch_size: int) -> list[tuple[str, list[str]]]:
        """Таблицы для модели частями по ``batch_size`` строк с ID.

        Возвращает ``[(markdown, [ID в этой части]), ...]``. В каждой части:
          - вводные строки листа (всё до первого ID: инструкции клиента);
          - несколько строк перед частью без ID — чтобы на границе частей не
            терялся заголовок раздела («Вирусные гепатиты:» над «В», «С»);
          - сами строки части; строки без ID (детализация) остаются с
            соседней строкой с ID над ними.
        """
        batch_size = max(1, batch_size)
        result: list[tuple[str, list[str]]] = []

        for table in template.tables:
            first_id = next(i for i, ln in enumerate(table.lines) if ln.row_id)
            preamble = table.lines[:first_id]

            groups: list[tuple[int, int]] = []     # [start, end) в table.lines
            start, ids_in_current = first_id, 0
            for i in range(first_id, len(table.lines)):
                if table.lines[i].row_id and ids_in_current == batch_size:
                    groups.append((start, i))
                    start, ids_in_current = i, 0
                if table.lines[i].row_id:
                    ids_in_current += 1
            groups.append((start, len(table.lines)))

            for start, end in groups:
                before = [
                    TableLine(row_id=None, cells=ln.cells)
                    for ln in table.lines[max(first_id, start - _CONTEXT_BEFORE):start]
                ]
                group = table.lines[start:end]
                md = TemplateReader._render(table, preamble + before + group)
                result.append((md, [ln.row_id for ln in group if ln.row_id]))
        return result

    @staticmethod
    def _render(table: SheetTable, lines: list[TableLine]) -> str:
        """Порядок столбцов: ID, «Требование», затем остальные как контекст.

        Модель склонна брать за требование первый столбец с текстом. Если там
        номер или название раздела (в Е-СОМ: B «2.46», C «первичные и повторные
        приемы специалистов», требование — в D), она повторяет не тот текст,
        и проверка отбрасывает всю часть.
        """
        req = table.columns.index(table.requirement_column)
        order = [req] + [i for i in range(len(table.columns)) if i != req]

        def title(i: int) -> str:
            letter = table.columns[i]
            name = table.titles.get(letter)
            return f"{letter}: {name}" if name else f"столбец {letter}"

        header = ["ID", f"Требование ({title(req)})"] + [
            f"Контекст ({title(i)})" for i in order[1:]
        ]
        out = [
            f"## Лист: {table.sheet}",
            "",
            "| " + " | ".join(header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
        ]
        for ln in lines:
            out.append("| " + " | ".join([ln.row_id or ""] + [ln.cells[i] for i in order]) + " |")
        return "\n".join(out)
