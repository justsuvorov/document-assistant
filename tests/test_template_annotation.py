"""Заполнение шаблона клиента (аннотация исходных листов).

Кейсы взяты из разбора «ОУ для сравнения.xlsx», где в шаблон легло 41 из 84
строк: затирался столбец «комментарии ВСК», одинаковые подзаголовки разделов
сваливались в первое вхождение, сокращения «…» и маркеры «·» не находились.
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from document_assistant.reports.report_models import InsuranceReport, ReportRow
from document_assistant.reports.writers import ExcelReportWriter

SHEET = "Объем услуг"


def _make_template(path: Path, rows: list[list]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _write(tmp_path: Path, template: Path, answers: list[tuple[str, str]]):
    report = InsuranceReport(
        rows=[ReportRow(req, "покрытие", status, f"коммент: {req[:20]}") for req, status in answers]
    )
    out = tmp_path / "out.xlsx"
    ExcelReportWriter().write(report, out, template)
    return load_workbook(out)[SHEET]


def _status_col(ws) -> int:
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col).value == "Статус":
            return col
    raise AssertionError("колонка 'Статус' не найдена")


class TestClientColumnsPreserved:
    def test_existing_second_column_is_not_overwritten(self, tmp_path):
        template = _make_template(tmp_path / "t.xlsx", [
            ["ОБЪЕМ УСЛУГ", "комментарии ВСК"],
            ["Лабораторная диагностика в сетевой лаборатории", "УВТ у нас 3 сеанса"],
        ])
        ws = _write(tmp_path, template, [("Лабораторная диагностика в сетевой лаборатории", "Есть")])

        assert ws["B1"].value == "комментарии ВСК"
        assert ws["B2"].value == "УВТ у нас 3 сеанса"

    def test_annotation_starts_after_last_used_column(self, tmp_path):
        template = _make_template(tmp_path / "t.xlsx", [
            ["Требование", "Комментарий клиента"],
            ["Стоматология", None],
        ])
        ws = _write(tmp_path, template, [("Стоматология", "Нет")])

        assert ws["C1"].value == "Покрытие по программе"
        assert ws["D1"].value == "Статус"
        assert ws["E1"].value == "Комментарий"
        assert ws["D2"].value == "Нет"


class TestRepeatedSectionHeaders:
    def test_same_header_in_different_sections_is_annotated_each_time(self, tmp_path):
        template = _make_template(tmp_path / "t.xlsx", [
            ["ОБЪЕМ УСЛУГ"],
            ["АМБУЛАТОРНАЯ ПОМОЩЬ"],
            ["Предоставляемые услуги:"],
            ["приемы врачей-специалистов"],
            ["СТОМАТОЛОГИЯ"],
            ["Предоставляемые услуги:"],
            ["лечение кариеса"],
        ])
        ws = _write(tmp_path, template, [
            ("АМБУЛАТОРНАЯ ПОМОЩЬ", "—"),
            ("Предоставляемые услуги:", "—"),
            ("приемы врачей-специалистов", "Есть"),
            ("СТОМАТОЛОГИЯ", "—"),
            ("Предоставляемые услуги:", "—"),
            ("лечение кариеса", "Частично"),
        ])
        col = _status_col(ws)

        assert ws.cell(row=3, column=col).value == "—"
        assert ws.cell(row=6, column=col).value == "—"
        assert ws.cell(row=7, column=col).value == "Частично"

    def test_similar_text_prefers_following_section(self, tmp_path):
        """Одинаковая фраза в двух разделах: ответ из второго раздела должен лечь
        во второй раздел, а не в первый."""
        template = _make_template(tmp_path / "t.xlsx", [
            ["ОБЪЕМ УСЛУГ"],
            ["ПОМОЩЬ НА ДОМУ"],
            ["Для застрахованных в Москве:"],
            ["СКОРАЯ ПОМОЩЬ"],
            ["Для застрахованных в Москве:"],
        ])
        ws = _write(tmp_path, template, [
            ("СКОРАЯ ПОМОЩЬ", "—"),
            ("Для застрахованных в Москве:", "Есть"),
        ])
        col = _status_col(ws)

        assert ws.cell(row=5, column=col).value == "Есть"
        assert ws.cell(row=3, column=col).value is None


class TestLlmTextVariants:
    def test_abbreviated_long_cell_is_matched(self, tmp_path):
        long_cell = (
            "первичный, повторный, консультативный приемы врачей-специалистов: "
            "аллерголога; гастроэнтеролога; гинеколога; дерматолога; кардиолога; "
            "невролога; нефролога; онколога; оториноларинголога; офтальмолога; "
            "проктолога; пульмонолога; ревматолога; терапевта; травматолога; уролога"
        )
        template = _make_template(tmp_path / "t.xlsx", [["ОБЪЕМ УСЛУГ"], [long_cell]])
        ws = _write(tmp_path, template, [
            ("первичный, повторный, консультативный приемы врачей-специалистов: аллерголога; ...", "Есть"),
        ])

        assert ws.cell(row=2, column=_status_col(ws)).value == "Есть"

    def test_unicode_ellipsis_and_repeated_ellipsis(self, tmp_path):
        long_cell = (
            "лечебно-оздоровительные процедуры и мероприятия: аутогемотерапия "
            "(1 курс - 10 сеансов); грязелечение; жемчужные ванны; ударно-волновая "
            "терапия; гирудотерапия; озонотерапия; карбокситерапия; прессотерапия"
        )
        template = _make_template(tmp_path / "t.xlsx", [["ОБЪЕМ УСЛУГ"], [long_cell]])
        ws = _write(tmp_path, template, [
            ("лечебно-оздоровительные процедуры и мероприятия: аутогемотерапия ...; …", "Частично"),
        ])

        assert ws.cell(row=2, column=_status_col(ws)).value == "Частично"

    def test_bullet_with_nbsp_in_template_is_matched(self, tmp_path):
        template = _make_template(tmp_path / "t.xlsx", [
            ["ОБЪЕМ УСЛУГ"],
            ["·" + "\xa0" * 10 + " лабораторная диагностика; "],
        ])
        ws = _write(tmp_path, template, [("лабораторная диагностика;", "Есть")])

        assert ws.cell(row=2, column=_status_col(ws)).value == "Есть"

    def test_unrelated_text_is_not_matched(self, tmp_path):
        template = _make_template(tmp_path / "t.xlsx", [
            ["ОБЪЕМ УСЛУГ"],
            ["Застрахованный может сдать анализы в сетевой лаборатории"],
        ])
        ws = _write(tmp_path, template, [("С записью в лабораторную сетевую лабораторию.", "Есть")])

        # Ничего не сопоставилось — лист клиента остаётся нетронутым,
        # ответ при этом есть на листе «Результат».
        assert ws.max_column == 1


def _hidden_category_template(path: Path) -> Path:
    """Как «ОУ.xlsx»: в скрытом A повторяется категория, требования — в D."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([None, None, None, "ОБЪЕМ УСЛУГ", "ДЕЙСТВУЮЩИЕ УСЛОВИЯ"])
    ws.append(["ВИД ПОМОЩИ", "БЛОК", "ДЕТАЛИЗАЦИЯ", "(!) ОБЪЕМ ПРИМЕНИМ", "подтверждение"])
    ws.append(["Страховым случаем является", None, None, "остром заболевании (состоянии)", None])
    ws.append(["Страховым случаем является", None, None, "обострении хронического заболевания", None])
    ws.append(["Страховым случаем является", None, None, "несчастном случае, в том числе травме", None])
    for col in "ABC":
        ws.column_dimensions[col].hidden = True
    wb.save(path)
    return path


class TestHiddenCategoryColumns:
    def test_requirement_column_is_detected_not_assumed_a(self, tmp_path):
        template = _hidden_category_template(tmp_path / "t.xlsx")
        ws = _write(tmp_path, template, [
            ("остром заболевании (состоянии)", "Есть"),
            ("обострении хронического заболевания", "Есть"),
            ("несчастном случае, в том числе травме", "Частично"),
        ])
        col = _status_col(ws)

        assert [ws.cell(row=r, column=col).value for r in (3, 4, 5)] == ["Есть", "Есть", "Частично"]

    def test_hidden_and_client_columns_untouched(self, tmp_path):
        template = _hidden_category_template(tmp_path / "t.xlsx")
        ws = _write(tmp_path, template, [("остром заболевании (состоянии)", "Есть")])

        assert ws["A3"].value == "Страховым случаем является"
        assert ws["D3"].value == "остром заболевании (состоянии)"
        assert ws["E2"].value == "подтверждение"
        assert _status_col(ws) > 5

    def test_hidden_columns_are_not_sent_to_model(self, tmp_path):
        from document_assistant.core.parsers import DataParser

        template = _hidden_category_template(tmp_path / "t.xlsx")
        text = DataParser(str(template)).origin_data(str(template))

        assert "остром заболевании" in text
        assert "Страховым случаем является" not in text
        assert "ВИД ПОМОЩИ" not in text
