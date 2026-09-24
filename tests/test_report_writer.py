from pathlib import Path

from openpyxl import load_workbook

from document_assistant.reports.report_models import InsuranceReport, ReportRow
from document_assistant.reports.writers import ExcelReportWriter


def _sample_report() -> InsuranceReport:
    return InsuranceReport(
        rows=[
            ReportRow("ДТП", "Полное покрытие", "Есть", "Лимит 1 000 000"),
            ReportRow("Пожар", "Не покрывается", "Нет", "Требуется доп. полис"),
        ],
        summary="Программа покрывает 1 из 2 требований.",
    )


class TestExcelReportWriterRawAnswerSheet:
    """Аннотация в шаблон клиента — эвристическая (сопоставление по тексту) и на
    плохих/испорченных шаблонах может промахнуться. Лист "Результат" — гарантия,
    что полный корректный ответ есть всегда, независимо от качества шаблона.
    """

    def test_adds_raw_answer_sheet(self, excel_file: Path, tmp_path: Path):
        output_path = tmp_path / "test_ответ.xlsx"
        ExcelReportWriter().write(_sample_report(), output_path, excel_file)

        wb = load_workbook(output_path)
        assert "Результат" in wb.sheetnames

    def test_raw_answer_sheet_is_active(self, excel_file: Path, tmp_path: Path):
        output_path = tmp_path / "test_ответ.xlsx"
        ExcelReportWriter().write(_sample_report(), output_path, excel_file)

        wb = load_workbook(output_path)
        assert wb.active.title == "Результат"

    def test_raw_answer_sheet_has_all_rows_in_standard_format(self, excel_file: Path, tmp_path: Path):
        output_path = tmp_path / "test_ответ.xlsx"
        ExcelReportWriter().write(_sample_report(), output_path, excel_file)

        wb = load_workbook(output_path)
        ws = wb["Результат"]
        assert [c.value for c in ws[1]] == [
            "Требование клиента", "Покрытие по программе", "Статус", "Комментарий",
        ]
        assert [c.value for c in ws[2]] == ["ДТП", "Полное покрытие", "Есть", "Лимит 1 000 000"]
        assert [c.value for c in ws[3]] == ["Пожар", "Не покрывается", "Нет", "Требуется доп. полис"]

    def test_original_template_sheets_preserved(self, excel_file: Path, tmp_path: Path):
        """Лист "Результат" добавляется, а не заменяет собой шаблон клиента."""
        output_path = tmp_path / "test_ответ.xlsx"
        ExcelReportWriter().write(_sample_report(), output_path, excel_file)

        wb = load_workbook(output_path)
        assert "Данные" in wb.sheetnames

    def test_title_deduplicated_if_client_sheet_already_named_result(self, tmp_path: Path):
        """Если в шаблоне клиента уже есть лист "Результат", наш не должен его затереть."""
        from openpyxl import Workbook

        src = tmp_path / "clash.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Результат"
        ws.append(["Что-то своё"])
        wb.save(src)

        output_path = tmp_path / "clash_ответ.xlsx"
        ExcelReportWriter().write(_sample_report(), output_path, src)

        out_wb = load_workbook(output_path)
        assert out_wb.sheetnames.count("Результат") == 0 or "Результат (2)" in out_wb.sheetnames
        assert "Результат (2)" in out_wb.sheetnames
        assert out_wb["Результат"]["A1"].value == "Что-то своё"
