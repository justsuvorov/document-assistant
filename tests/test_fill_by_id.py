"""Заполнение шаблона по ID строк (document_assistant/filling).

Главное свойство: ответ не может лечь в чужую строку. Всё, в чём нет
уверенности, остаётся «Не обработано».
"""

import re
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from document_assistant.ai.preprocessor import ProcessingTask
from document_assistant.filling.answers import AnswerParser, AnswerValidator, text_agrees
from document_assistant.filling.models import RowAnswer
from document_assistant.filling.prompt import RESPONSE_CONTRACT, FillPromptBuilder
from document_assistant.filling.reader import TemplateReader
from document_assistant.filling.service import TemplateFillService
from document_assistant.filling.writer import NOT_PROCESSED

SHEET = "Объем услуг"
REQS = [
    "Амбулаторно-поликлиническая помощь",
    "приемы врачей-специалистов: терапевт, невролог",
    "курс лечения у остеопата",
    "PRP-терапия при патологии опорно-двигательного аппарата",
    "аутогемотерапия",
]


def _template(path: Path, extra=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(["Требование", "комментарии клиента"])
    for r in REQS:
        ws.append([r, None])
    if extra:
        extra(wb)
    wb.save(path)
    return path


def _rows_in_prompt(prompt: str) -> list[tuple[str, str]]:
    """(ID, первая ячейка после ID) из таблицы запроса в промпте."""
    out = []
    for line in prompt.splitlines():
        m = re.match(r"^\| (R\d{4}) \| ([^|]*) \|", line)
        if m:
            out.append((m.group(1), m.group(2).strip()))
    return out


class FakeModel:
    """Отвечает по ID из промпта. ``tamper(id, text) -> list[row] | None`` —
    подмена строк ответа для имитации сбоев модели."""

    def __init__(self, tamper=None):
        self.prompts = []
        self._tamper = tamper

    def response(self, prompt: str) -> str:
        self.prompts.append(prompt)
        lines = ["| ID | Требование (начало) | Покрытие по программе | Статус | Комментарий |",
                 "|---|---|---|---|---|"]
        for rid, text in _rows_in_prompt(prompt):
            rows = [[rid, text, f"покрытие {rid}", "Есть", f"коммент {rid}"]]
            if self._tamper:
                rows = self._tamper(rid, text, rows)
            lines += ["| " + " | ".join(r) + " |" for r in rows]
        return "\n".join(lines) + "\n\n## Резюме\nПрограмма Комплексная."


def _run(tmp_path, template, model, batch_size=50, monkeypatch=None):
    if monkeypatch:
        from document_assistant.core.settings import settings
        monkeypatch.setattr(settings, "llm_batch_size", batch_size)
    task = ProcessingTask(request_id=1, file_path=str(template))
    service = TemplateFillService(
        task=task,
        ai_model=model,
        prompt_builder=FillPromptBuilder("роль", "", str(tmp_path / "нет нормативки"), 400_000),
        fallback=lambda: pytest.fail("не должен уходить в старый путь"),
    )
    out = service.result()
    return load_workbook(out["output_file"])


def _status(ws, row: int) -> str:
    for c in range(1, ws.max_column + 1):
        if ws.cell(1, c).value == "Статус":
            return ws.cell(row, c).value
    raise AssertionError("нет колонки Статус")


# ── Сквозные сценарии ─────────────────────────────────────────────────────────

class TestEndToEnd:
    def test_all_rows_filled_by_id(self, tmp_path):
        wb = _run(tmp_path, _template(tmp_path / "t.xlsx"), FakeModel())
        ws = wb[SHEET]
        assert [_status(ws, r) for r in range(2, 7)] == ["Есть"] * 5
        assert ws["B1"].value == "комментарии клиента"   # колонки клиента не тронуты
        assert [ws.cell(1, c).value for c in range(3, 7)] == [
            "Покрытие по программе", "Статус", "Комментарий", "ID"]
        assert ws.max_column == 6                         # ID — самый правый столбец
        assert ws["F4"].value == "R0003"                  # R0003 → 3-е требование → строка 4
        assert ws["C4"].value == "покрытие R0003"

    def test_wrong_text_is_not_written(self, tmp_path):
        """Модель указала ID остеопата, но текст от другой строки — не пишем."""
        def tamper(rid, text, rows):
            if rid == "R0003":
                rows[0][1] = "аллерген-специфическая иммунотерапия"
            return rows
        ws = _run(tmp_path, _template(tmp_path / "t.xlsx"), FakeModel(tamper))[SHEET]
        assert _status(ws, 4) == NOT_PROCESSED
        assert _status(ws, 5) == "Есть"

    def test_duplicate_id_keeps_first(self, tmp_path):
        def tamper(rid, text, rows):
            if rid == "R0002":
                rows.append([rid, text, "ДУБЛЬ", "Нет", "дубль"])
            return rows
        ws = _run(tmp_path, _template(tmp_path / "t.xlsx"), FakeModel(tamper))[SHEET]
        assert ws["C3"].value == "покрытие R0002"

    def test_missing_answer_marked_not_processed(self, tmp_path):
        ws = _run(tmp_path, _template(tmp_path / "t.xlsx"),
                  FakeModel(lambda rid, text, rows: [] if rid == "R0004" else rows))[SHEET]
        assert _status(ws, 5) == NOT_PROCESSED
        assert ws["F5"].value == "R0004"                  # ID есть и у необработанной строки
        assert ws["C5"].value in (None, "")

    def test_answer_for_id_from_other_chunk_rejected(self, tmp_path, monkeypatch):
        """В части 2 модель вернула ID из части 1 — отбрасываем, не перезаписываем."""
        def tamper(rid, text, rows):
            if rid == "R0004":
                rows.append(["R0001", REQS[0], "ЧУЖОЕ", "Нет", "чужое"])
            return rows
        ws = _run(tmp_path, _template(tmp_path / "t.xlsx"), FakeModel(tamper),
                  batch_size=3, monkeypatch=monkeypatch)[SHEET]
        assert ws["C2"].value == "покрытие R0001"

    def test_result_sheet_lists_every_row_with_address(self, tmp_path):
        wb = _run(tmp_path, _template(tmp_path / "t.xlsx"),
                  FakeModel(lambda rid, text, rows: [] if rid == "R0005" else rows))
        assert wb.sheetnames[0] == "Результат"
        res = wb["Результат"]
        assert [c.value for c in res[1]][:4] == ["ID", "Лист", "Строка", "Требование клиента"]
        assert [res.cell(r, 1).value for r in range(2, 7)] == [f"R000{i}" for i in range(1, 6)]
        assert res.cell(4, 3).value == 4
        assert res.cell(6, 6).value == NOT_PROCESSED
        assert "Резюме" in wb.sheetnames

    def test_chunks_cover_all_rows(self, tmp_path, monkeypatch):
        model = FakeModel()
        ws = _run(tmp_path, _template(tmp_path / "t.xlsx"), model, batch_size=2,
                  monkeypatch=monkeypatch)[SHEET]
        assert len(model.prompts) == 3
        assert [_status(ws, r) for r in range(2, 7)] == ["Есть"] * 5
        # шапка листа повторяется в каждой части
        assert all("| Требование | комментарии клиента |" in p for p in model.prompts)

    def test_prompt_has_contract_and_rows_last(self, tmp_path):
        model = FakeModel()
        _run(tmp_path, _template(tmp_path / "t.xlsx"), model)
        p = model.prompts[0]
        assert RESPONSE_CONTRACT.strip() in p
        assert p.index("| R0001 |") < p.index("## ФОРМАТ ОТВЕТА")

    def test_empty_template_goes_to_fallback(self, tmp_path):
        wb = Workbook()
        wb.active.append(["только шапка"])
        path = tmp_path / "e.xlsx"
        wb.save(path)

        class Legacy:
            def result(self, max_chunks_override=0):
                return {"output_file": "legacy"}

        service = TemplateFillService(
            task=ProcessingTask(request_id=1, file_path=str(path)),
            ai_model=FakeModel(),
            prompt_builder=FillPromptBuilder("роль", "", str(tmp_path / "нет"), 400_000),
            fallback=Legacy,
        )
        assert service.result() == {"output_file": "legacy"}


# ── Чтение шаблона ────────────────────────────────────────────────────────────

class TestReader:
    def test_hidden_sheet_and_column_skipped_rows_kept(self, tmp_path):
        def extra(wb):
            ws = wb[SHEET]
            ws.insert_cols(1)
            for r in range(1, ws.max_row + 1):
                ws.cell(r, 1, "Категория повторяется в каждой строке")
            ws.column_dimensions["A"].hidden = True
            ws.row_dimensions[5].hidden = True        # скрыта фильтром — не пропускаем
            clinics = wb.create_sheet("Клиники")
            for i in range(30):
                clinics.append([f"Клиника номер {i}", "Москва"])
            clinics.sheet_state = "hidden"

        t = TemplateReader().read(_template(tmp_path / "t.xlsx", extra))
        assert [r.text for r in t.rows] == REQS
        assert {r.sheet for r in t.rows} == {SHEET}
        md = TemplateReader.chunks(t, 50)[0][0]
        assert "Категория повторяется" not in md
        assert "Клиника" not in md

    def test_detail_rows_without_id_stay_with_parent(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["Услуга", "Детализация"])
        ws.append(["Анализ крови", "клинический"])
        ws.append([None, "биохимический"])
        ws.append(["Электрокардиография", "в покое"])
        path = tmp_path / "d.xlsx"
        wb.save(path)

        t = TemplateReader().read(path)
        assert [r.text for r in t.rows] == ["Анализ крови", "Электрокардиография"]
        chunks = TemplateReader.chunks(t, 1)
        assert "биохимический" in chunks[0][0] and chunks[0][1] == ["R0001"]
        assert "биохимический" not in chunks[1][0]


# ── Разбор ответа и проверка ──────────────────────────────────────────────────

class TestParser:
    def test_variants(self):
        raw = """Вот таблица:
| ID | Требование (начало) | Покрытие | Статус | Комментарий |
|---|---|---|---|---|
| R0001 | Амбулаторно-поликлиническая | п1 | Есть | к1 |
| **R0002** | приемы врачей | п2 | Частично | к2 |
| R3 | курс лечения | п3 | Нет | к3 |
| R0004 | п4 | Есть | к4 |
| R0005 | аутогемотерапия | п5 | Есть | часть | вторая часть |
| Итого | | | | |

## Резюме
Программа А."""
        answers, summary = AnswerParser().parse(raw)
        assert [a.row_id for a in answers] == ["R0001", "R0002", "R0003", "R0004", "R0005"]
        assert answers[3].requirement == "" and answers[3].coverage == "п4"
        assert answers[4].comment == "часть | вторая часть"
        assert summary == "Программа А."

    def test_summary_table_not_parsed(self):
        raw = "| R0001 | а | б | Есть | в |\n\n## Резюме\n| R0002 | x | y | Нет | z |"
        assert [a.row_id for a in AnswerParser().parse(raw)[0]] == ["R0001"]


@pytest.mark.parametrize("model_text, template_text, ok", [
    ("курс лечения у остеопата", "курс лечения у остеопата", True),
    ("первичный, повторный, консультативный приемы...", "первичный, повторный, консультативный приемы врачей", True),
    ("Анализ крови", "Анализ крови", True),
    ("МРТ", "МРТ", True),
    ("ЭХО-КГ и ЭХО-ЭГ", "ЭХО -КГ и ЭХО-ЭГ", True),
    ("курс лечения остеопатом", "курс лечения у остеопата", True),
    ("", "что угодно", True),
    ("аллерген-специфическая иммунотерапия", "курс лечения у остеопата", False),
    ("Урология", "Эндокринология", False),
    ("лечение заболеваний вен: склеротерапия", "лечение с использованием аппаратов квантовой терапии", False),
    # модель дописала детализацию из соседнего столбца
    ("Анализ крови / - клинический развернутый", "Анализ крови", True),
    # заголовок раздела не сходит за начало следующей строки
    ("САХАРНЫЙ ДИАБЕТ", "Сахарный диабет 1 тип впервые выявленный", False),
    # строки различаются только числом
    ("Осложнения сахарного диабета 1 типа", "Осложнения сахарного диабета 2 типа", False),
    # различаются только первым словом при общем хвосте
    ("водолечение в лечебных целях", "грязелечение в лечебных целях", False),
    # «хирург» — не начало «хирургического»
    ("Хирург", "Хирургическое лечение пародонтита", False),
    # начало длинного текста (8+ слов) — принимается
    ("расходные материалы при нейрохирургических операциях в связи с травмой",
     "расходные материалы при нейрохирургических операциях в связи с травмой по плановым показаниям", True),
    # то же начало, но строка про заболевание, а не травму
    ("расходные материалы при нейрохирургических операциях в связи с травмой по",
     "расходные материалы при нейрохирургических операциях в связи заболеванием по плановым показаниям", False),
])
def test_text_agrees(model_text, template_text, ok):
    assert text_agrees(model_text, template_text) is ok


def test_validator_reasons(tmp_path):
    t = TemplateReader().read(_template(tmp_path / "t.xlsx"))
    v = AnswerValidator(t.by_id())
    ok, bad = v.add([
        RowAnswer("R0001", REQS[0], "", "Есть", ""),
        RowAnswer("R0001", REQS[0], "", "Нет", ""),
        RowAnswer("R0009", "x", "", "Есть", ""),
        RowAnswer("R0002", "совсем другой текст про стоматологию", "", "Есть", ""),
    ], allowed_ids={"R0001", "R0002"})
    assert (ok, bad) == (1, 3)
    assert [r.reason for r in v.rejected] == [
        "повтор ID", "ID не из этой части запроса", "текст не совпадает со строкой шаблона"]


# ── Оркестратор ───────────────────────────────────────────────────────────────

class TestFactory:
    @pytest.fixture(autouse=True)
    def _no_model(self, monkeypatch):
        from document_assistant.services import factory
        monkeypatch.setattr(factory.ModelFactory, "create", staticmethod(lambda: FakeModel()))

    def _build(self, monkeypatch, mode, path):
        from document_assistant.core.settings import settings
        from document_assistant.services.factory import build_dms_service
        monkeypatch.setattr(settings, "fill_mode", mode)
        return build_dms_service(ProcessingTask(request_id=1, file_path=str(path)),
                                 normative_base=str(path.parent / "нет"))

    def test_xlsx_by_id(self, monkeypatch, tmp_path):
        assert type(self._build(monkeypatch, "by_id", tmp_path / "a.xlsx")).__name__ == "TemplateFillService"

    def test_legacy_switch(self, monkeypatch, tmp_path):
        assert type(self._build(monkeypatch, "legacy", tmp_path / "a.xlsx")).__name__ == "AIAssistantService"

    def test_docx_always_legacy(self, monkeypatch, tmp_path):
        assert type(self._build(monkeypatch, "by_id", tmp_path / "a.docx")).__name__ == "AIAssistantService"
